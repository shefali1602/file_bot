from flask import Flask, render_template, request, jsonify
import os
from PyPDF2 import PdfReader
from docx import Document
from pptx import Presentation
import openpyxl
import subprocess
import sys

app = Flask(__name__)

SUPPORTED_EXT = [".pdf", ".docx", ".pptx", ".txt", ".xlsx"]
file_index = {}
folder_indexed = False
folder_path = ""

def extract_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    text = ""
    try:
        if ext == ".pdf":
            reader = PdfReader(filepath)
            text = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
        elif ext == ".docx":
            doc = Document(filepath)
            text = "\n".join([p.text for p in doc.paragraphs])
        elif ext == ".pptx":
            prs = Presentation(filepath)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        elif ext == ".txt":
            with open(filepath, "r", encoding="utf-8") as f:
                text = f.read()
        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(filepath, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    text += " ".join([str(cell) for cell in row if cell]) + "\n"
    except Exception as e:
        print(f"Failed to read {filepath}: {e}")
    return text.lower()

def index_files(root_dir):
    index = {}
    for dirpath, _, filenames in os.walk(root_dir):
        for file in filenames:
            filepath = os.path.join(dirpath, file)
            if os.path.splitext(file)[1].lower() in SUPPORTED_EXT:
                print(f"Indexing: {filepath}")
                index[filepath] = extract_text(filepath)
    return index

def search_query(query, index):
    query = query.lower()
    results = []
    for path, content in index.items():
        if query in content:
            results.append(path)
    return results

def open_file(filepath):
    try:
        if os.name == 'nt':
            os.startfile(filepath)
        elif os.name == 'posix':
            subprocess.call(['open' if sys.platform == 'darwin' else 'xdg-open', filepath])
    except Exception as e:
        print(f"Error opening file: {e}")

@app.route("/")
def home():
    return render_template("chatbot.html")

@app.route("/message", methods=["POST"])
def message():
    global folder_indexed, folder_path, file_index

    data = request.get_json()
    user_message = data.get("message", "").strip()

    if not user_message:
        return jsonify({"response": {"type": "text", "text": "Hi! I'm Inquiro. Please provide a folder path to begin."}})

    if os.path.isdir(user_message):
        folder_path = user_message
        file_index = index_files(folder_path)
        folder_indexed = True

        file_list = [{"filename": os.path.basename(f), "filepath": f} for f in file_index.keys()]
        return jsonify({
            "response": {
                "type": "file_list",
                "files": file_list,
                "count": len(file_list)
            }
        })

    if folder_indexed:
        matches = search_query(user_message, file_index)

        if matches:
            file_buttons = [{"filename": os.path.basename(f), "filepath": f} for f in matches]
            return jsonify({"response": {"type": "files", "files": file_buttons}})
        else:
            return jsonify({"response": {"type": "text", "text": "No matching files found."}})

    return jsonify({"response": {"type": "text", "text": "Please provide a valid folder path first."}})

@app.route("/open_file", methods=["POST"])
def open_file_route():
    file_path = request.json.get("file_path")
    try:
        if os.path.isfile(file_path):
            open_file(file_path)
            return jsonify({"message": f"Opening file: {os.path.basename(file_path)}"})
        else:
            return jsonify({"message": ""
            "File not found."})
    except Exception as e:
        return jsonify({"message": f"Error opening file: {e}"})

if __name__ == "__main__":
    app.run(debug=True)
